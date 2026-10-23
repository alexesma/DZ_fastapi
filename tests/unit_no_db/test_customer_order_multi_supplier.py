from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from dz_fastapi.models.partner import CustomerOrderItem, SupplierOrderItem
from dz_fastapi.services import customer_orders as service
from dz_fastapi.services.partssoft_order_reconciliation import (
    local_order_fingerprint,
    remote_order_fingerprint,
)


def _offer(provider_id, quantity, purchase_price, customer_offer_price):
    return service.OfferRow(
        autopart_id=4680,
        provider_id=provider_id,
        provider_config_id=provider_id + 100,
        quantity=quantity,
        price=customer_offer_price,
        supplier_price=purchase_price,
        is_own_price=False,
        vat_rate=22,
    )


def test_one_row_uses_cheaper_supplier_then_fills_from_second_supplier():
    aruda = _offer(930, 78, 32.9, 82.25)
    vladimirberg = _offer(938, 156, 80, 88)
    aruda.alternatives = (vladimirberg,)
    config = SimpleNamespace(price_warning_pct=5)
    row = service.ParsedOrderRow(6, "9030122013", "TOYOTA", None, 150, 88)
    allocated = {}

    plan = service._plan_customer_order_allocations(config, row, aruda, None, allocated)

    assert [(part.offer.provider_id, part.quantity) for part in plan.allocations] == [
        (930, 78), (938, 72),
    ]
    assert plan.reject_qty == 0
    assert allocated == {(930, 4680): 78, (938, 4680): 72}

    second_row = service.ParsedOrderRow(7, "9030122013", "TOYOTA", None, 100, 88)
    second_plan = service._plan_customer_order_allocations(
        config, second_row, aruda, None, allocated,
    )
    assert [(part.offer.provider_id, part.quantity) for part in second_plan.allocations] == [
        (938, 84),
    ]
    assert second_plan.reject_qty == 16


def test_expensive_candidate_is_skipped_before_allocating_stock():
    expensive = _offer(930, 78, 32.9, 120)
    acceptable = _offer(938, 156, 80, 88)
    expensive.alternatives = (acceptable,)
    row = service.ParsedOrderRow(6, "9030122013", "TOYOTA", None, 72, 88)
    allocated = {}

    plan = service._plan_customer_order_allocations(
        SimpleNamespace(price_warning_pct=5), row, expensive, None, allocated,
    )

    assert [(part.offer.provider_id, part.quantity) for part in plan.allocations] == [
        (938, 72),
    ]
    assert plan.reject_qty == 0
    assert (930, 4680) not in allocated


def test_profit_order_accounts_for_input_vat():
    with_vat = _offer(930, 5, 120, 130)
    without_vat = _offer(938, 5, 105, 130)
    without_vat.vat_rate = None

    assert sorted((without_vat, with_vat), key=service._offer_profit_rank) == [
        with_vat, without_vat,
    ]


@pytest.mark.asyncio
async def test_current_offers_keep_second_supplier_and_rank_by_net_cost(monkeypatch):
    def source(provider_id, config_id, markup):
        provider = SimpleNamespace(vat_rate=22, is_vat_payer=True)
        provider_config = SimpleNamespace(
            provider=provider, mailing_included_brands=[],
        )
        return SimpleNamespace(
            enabled=True,
            provider_config_id=config_id,
            provider_config=provider_config,
            markup=markup,
            brand_markups={},
            brand_filters={},
            position_filters={},
            min_price=None,
            max_price=None,
            min_quantity=None,
            max_quantity=None,
            additional_filters={},
            provider_id=provider_id,
        )

    sources = [source(930, 33, 150), source(938, 44, 10)]
    config = SimpleNamespace(
        id=11,
        individual_markups={},
        default_filters={},
        brand_filters=[],
        category_filter=[],
        price_intervals=[],
        position_filters=[],
        supplier_quantity_filters=[],
        additional_filters={},
        own_filters={},
        other_filters={},
        supplier_filters={},
        general_markup=1,
        own_price_list_markup=1,
        third_party_markup=1,
    )
    monkeypatch.setattr(service, "_get_order_offer_sources", AsyncMock(
        return_value=sources,
    ))
    monkeypatch.setattr(
        service.crud_pricelist,
        "get_latest_pricelist_by_config",
        AsyncMock(
            side_effect=lambda session, provider_config_id: SimpleNamespace(
                id=provider_config_id,
            ),
        ),
    )

    async def fetch(_pricelist_id, _session, *, oem_numbers):
        provider_id, price, quantity = (
            (930, 32.9, 78) if _pricelist_id == 33 else (938, 80, 156)
        )
        assert oem_numbers == {"9030122013"}
        return pd.DataFrame([{
            "autopart_id": 4680,
            "provider_id": provider_id,
            "provider_config_id": _pricelist_id,
            "oem_number": "9030122013",
            "brand": "TOYOTA",
            "quantity": quantity,
            "price": price,
            "is_own_price": False,
        }])

    monkeypatch.setattr(service.crud_pricelist, "fetch_pricelist_dataframe", fetch)

    offers = await service._build_current_offers(
        None, config, required_oems={"9030122013"},
    )

    primary = offers[("9030122013", "TOYOTA")]
    assert primary.provider_id == 930
    assert primary.quantity == 78
    assert primary.price == pytest.approx(82.25)
    assert [(offer.provider_id, offer.quantity, offer.price)
            for offer in primary.alternatives] == [(938, 156, 88.0)]


@pytest.mark.asyncio
async def test_processing_creates_separate_supplier_lines_for_one_customer_row(monkeypatch):
    aruda = _offer(930, 78, 32.9, 82.25)
    aruda.alternatives = (_offer(938, 156, 80, 88),)
    monkeypatch.setattr(
        service,
        "_prepare_customer_order_context",
        AsyncMock(return_value=(
            {}, {("9030122013", "TOYOTA"): aruda}, {}, object(),
        )),
    )
    monkeypatch.setattr(service, "assert_customer_credit_available", AsyncMock())
    monkeypatch.setattr(service, "_get_or_create_open_supplier_order", AsyncMock(
        side_effect=lambda session, provider_id, provider_config_id: SimpleNamespace(
            id=provider_id + 1000,
        ),
    ))

    class Session:
        def __init__(self):
            self.added = []
            self.next_id = 1

        async def scalars(self, _statement):
            return SimpleNamespace(all=lambda: [])

        def add(self, item):
            if isinstance(item, CustomerOrderItem) and item.id is None:
                item.id = self.next_id
                self.next_id += 1
            self.added.append(item)

        async def flush(self):
            pass

        async def commit(self):
            pass

    session = Session()
    config = SimpleNamespace(customer_id=946, price_warning_pct=5, price_tolerance_pct=2)
    order = SimpleNamespace(id=5899, external_payload={})
    row = service.ParsedOrderRow(6, "9030122013", "TOYOTA", None, 150, 88)

    items, rejected = await service._process_manual_rows(
        session, config, order, [row],
    )

    assert rejected == []
    assert [(item.supplier_id, item.requested_qty, item.ship_qty) for item in items] == [
        (930, 78, 78), (938, 72, 72),
    ]
    supplier_lines = [item for item in session.added if isinstance(item, SupplierOrderItem)]
    assert [(item.supplier_order_id, item.quantity, float(item.price))
            for item in supplier_lines] == [(1930, 78, 32.9), (1938, 72, 80.0)]
    assert [item.customer_order_item_id for item in supplier_lines] == [1, 2]


def test_response_file_aggregates_allocations_into_original_row():
    items = [
        SimpleNamespace(row_index=2, ship_qty=78, reject_qty=0, requested_price=88),
        SimpleNamespace(row_index=2, ship_qty=72, reject_qty=0, requested_price=88),
    ]
    config = SimpleNamespace(
        qty_col=2,
        ship_qty_col=3,
        reject_qty_col=4,
        ship_price_col=5,
        ship_mode=service.CUSTOMER_ORDER_SHIP_MODE.WRITE_SHIP_QTY,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(row=2, column=3).value = 150
    source = BytesIO()
    workbook.save(source)
    source.seek(0)

    response = service._apply_response_updates_excel(source, config, items)
    updated = load_workbook(response).active

    assert updated.cell(row=2, column=4).value == 150
    assert updated.cell(row=2, column=6).value == 88

    csv_config = SimpleNamespace(**{
        **vars(config),
        "ship_mode": service.CUSTOMER_ORDER_SHIP_MODE.WRITE_REJECT_QTY,
    })
    csv_source = BytesIO(b"header,brand,qty,ship,reject,price\npart,TOYOTA,150,0,0,0\n")
    csv_items = [
        SimpleNamespace(row_index=1, ship_qty=78, reject_qty=0, requested_price=88),
        SimpleNamespace(row_index=1, ship_qty=67, reject_qty=0, requested_price=88),
        SimpleNamespace(row_index=1, ship_qty=0, reject_qty=5, requested_price=88),
    ]
    csv_response = service._apply_response_updates_csv(csv_source, csv_config, csv_items)
    csv_response.seek(0)
    csv = pd.read_csv(csv_response, header=None)
    assert int(csv.iat[1, 4]) == 5


def test_customer_forward_keeps_one_original_line_after_supplier_split():
    parts = [
        SimpleNamespace(
            row_index=6, oem="9030122013", brand="TOYOTA", name="Кольцо",
            requested_qty=78, requested_price=88, reject_reason_text=None,
        ),
        SimpleNamespace(
            row_index=6, oem="9030122013", brand="TOYOTA", name="Кольцо",
            requested_qty=72, requested_price=88, reject_reason_text=None,
        ),
    ]
    order = SimpleNamespace(id=5899, order_number="#", items=parts)

    rows, quantity, amount = service._build_customer_order_forward_rows(order)

    assert quantity == 150
    assert amount == 13200
    assert len(rows) == 1
    assert rows[0]["qty"] == 150
    assert rows[0]["sum"] == 13200


def test_partssoft_duplicate_matching_sees_original_customer_quantity():
    local = SimpleNamespace(items=[
        SimpleNamespace(
            row_index=6, oem="9030122013", brand="TOYOTA",
            requested_qty=78, requested_price=88,
        ),
        SimpleNamespace(
            row_index=6, oem="9030122013", brand="TOYOTA",
            requested_qty=72, requested_price=88,
        ),
    ])
    remote = {"order_items": [{
        "oem": "9030122013", "make_name": "TOYOTA", "qnt": 150, "cost": 88,
    }]}

    assert local_order_fingerprint(local) == remote_order_fingerprint(remote)
